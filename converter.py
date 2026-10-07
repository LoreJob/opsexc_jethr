"""Convert the five provider exports into Jet HR's fixed-width payroll format."""

from __future__ import annotations

import csv
import io
import re
import unicodedata
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterable

import xlrd
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
CONFIG_DIR = ROOT / "Kit Candidato" / "config"
PROVIDER_EXTENSIONS = {"A": ".xlsx", "B": ".csv", "C": ".xls", "D": ".csv", "E": ".xlsx"}


class ConversionError(ValueError):
    """The input cannot be converted at all."""


@dataclass(frozen=True)
class Movement:
    row: int
    fiscal_code: str
    display_name: str
    given_name: str
    surname: str
    treatment: str
    subcategory: str
    amount: object
    date_value: object = None


@dataclass(frozen=True)
class Issue:
    row: int
    person: str
    detail: str

    def as_dict(self) -> dict:
        return {"row": self.row, "person": self.person, "detail": self.detail}


@dataclass(frozen=True)
class ConversionResult:
    filename: str
    content: str
    input_rows: int
    converted_rows: int
    output_rows: int
    issues: list[Issue]


def _key(value: object) -> str:
    value = unicodedata.normalize("NFKC", str(value or ""))
    return " ".join(value.casefold().split())


def _tax_key(value: object) -> str:
    """Normalize spelling variants without erasing provider-specific context."""
    value = unicodedata.normalize("NFKD", str(value or "")).casefold()
    value = "".join(char for char in value if not unicodedata.combining(char))
    value = re.sub(r"\bart\s*\.?\s*(\d+)", r"art\1", value)
    value = re.sub(r"\bc\s*\.?\s*(\d+)", r"c\1", value)
    value = re.sub(r"\b(?:lettera|lett|let)\b", "lett", value)
    value = re.sub(r"\bf\s*[-.\s]*\s*b(?:is)?\b", "fbis", value)
    value = re.sub(r"\bd\s*[-.\s]*\s*b(?:is)?\b", "dbis", value)
    value = re.sub(r"\bf\s*[-.\s]*\s*ter\b", "fter", value)
    return " ".join(re.findall(r"[a-z0-9]+", value))


def _name_key(value: object) -> str:
    value = unicodedata.normalize("NFKD", str(value or ""))
    value = "".join(char for char in value if not unicodedata.combining(char))
    return " ".join(re.findall(r"[a-z0-9]+", value.casefold()))


def _text(value: object) -> str:
    return str(value).strip() if value is not None else ""


def _date_period(value: object) -> str:
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y%m")
    raw = _text(value)
    for date_format in ("%Y-%m-%d", "%Y%m%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(raw, date_format).strftime("%Y%m")
        except ValueError:
            continue
    raise ValueError(f"data mancante o non valida: {raw or '(vuota)'}")


def _validate_period(movements: list[Movement], selected_period: str) -> None:
    wrong_periods: dict[str, list[int]] = defaultdict(list)
    invalid_rows: list[int] = []
    for movement in movements:
        try:
            movement_period = _date_period(movement.date_value)
        except ValueError:
            invalid_rows.append(movement.row)
            continue
        if movement_period != selected_period:
            wrong_periods[movement_period].append(movement.row)
    if not wrong_periods and not invalid_rows:
        return

    details = []
    for period, rows in sorted(wrong_periods.items()):
        examples = ", ".join(str(row) for row in rows[:5])
        suffix = ", …" if len(rows) > 5 else ""
        details.append(f"{period}: {len(rows)} righe (es. {examples}{suffix})")
    if invalid_rows:
        examples = ", ".join(str(row) for row in invalid_rows[:5])
        suffix = ", …" if len(invalid_rows) > 5 else ""
        details.append(f"date mancanti o non valide: {len(invalid_rows)} righe (es. {examples}{suffix})")
    raise ConversionError(
        f"Data non coerente con il periodo selezionato ({selected_period}). "
        + "; ".join(details)
        + ". Seleziona il periodo corretto o carica un altro file."
    )


def _decode_csv(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        # Provider B's example is Windows-1252 encoded.
        return data.decode("cp1252")


def _headered_rows(rows: Iterable[tuple[int, tuple]], required: set[str]):
    header = None
    for row_number, values in rows:
        cells = tuple(_text(value) for value in values)
        if header is None:
            if required.issubset(set(cells)):
                header = {name: index for index, name in enumerate(cells)}
            continue
        if not any(cells):
            continue
        yield row_number, {name: values[index] if index < len(values) else None for name, index in header.items()}
    if header is None:
        raise ConversionError("Intestazioni attese non trovate nel file selezionato.")


def _xlsx_rows(data: bytes, required: set[str]):
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for sheet in workbook.worksheets:
            numbered = ((i, row) for i, row in enumerate(sheet.iter_rows(values_only=True), start=1))
            try:
                yield from _headered_rows(numbered, required)
                return
            except ConversionError:
                continue
    except Exception as exc:
        raise ConversionError("Impossibile leggere il file XLSX.") from exc
    raise ConversionError("Intestazioni attese non trovate in nessun foglio XLSX.")


def _xls_rows(data: bytes, required: set[str]):
    try:
        workbook = xlrd.open_workbook(file_contents=data)
        for sheet in workbook.sheets():
            numbered = (
                (
                    i + 1,
                    tuple(
                        xlrd.xldate_as_datetime(cell.value, workbook.datemode)
                        if cell.ctype == xlrd.XL_CELL_DATE else cell.value
                        for cell in sheet.row(i)
                    ),
                )
                for i in range(sheet.nrows)
            )
            try:
                yield from _headered_rows(numbered, required)
                return
            except ConversionError:
                continue
    except Exception as exc:
        raise ConversionError("Impossibile leggere il file XLS.") from exc
    raise ConversionError("Intestazioni attese non trovate in nessun foglio XLS.")


def _csv_rows(data: bytes, required: set[str]):
    reader = csv.reader(io.StringIO(_decode_csv(data)), delimiter=";")
    return _headered_rows(((i, tuple(row)) for i, row in enumerate(reader, start=1)), required)


def _parse_provider(provider: str, data: bytes) -> list[Movement]:
    movements = []
    if provider == "A":
        rows = _xlsx_rows(data, {"Codice fiscale dipendente", "Nome dipendente", "Tratt. Fiscale", "Importo", "Data"})
        for number, row in rows:
            movements.append(Movement(number, _text(row["Codice fiscale dipendente"]), _text(row["Nome dipendente"]), "", "", _text(row["Tratt. Fiscale"]), "", row["Importo"], row["Data"]))
    elif provider == "B":
        reader = csv.reader(io.StringIO(_decode_csv(data)), delimiter=";")
        for number, row in enumerate(reader, start=1):
            if not any(cell.strip() for cell in row):
                continue
            if len(row) < 9:
                movements.append(Movement(number, "", "", "", "", "", "", ""))
                continue
            surname, given = row[1].strip(), row[2].strip()
            movements.append(Movement(number, "", f"{given} {surname}".strip(), given, surname, row[3].strip(), "", row[8], row[4]))
    elif provider == "C":
        rows = _xls_rows(data, {"CF dipendente", "Art.", "Rif. normativi", "Segno a cedolino", "Importo totale", "Data movimento"})
        for number, row in rows:
            treatment = f"{_text(row['Art.'])} {_text(row['Rif. normativi'])}".strip()
            movements.append(Movement(number, _text(row["CF dipendente"]), f"{_text(row['Nome dipendente'])} {_text(row['Cognome dipendente'])}".strip(), "", "", treatment, _text(row["Segno a cedolino"]), row["Importo totale"], row["Data movimento"]))
    elif provider == "D":
        rows = _csv_rows(data, {"Nome dipendente", "Cognome dipendente", "Riferimento al TUIR", "Valore", "Data di assegnazione"})
        for number, row in rows:
            given, surname = _text(row["Nome dipendente"]), _text(row["Cognome dipendente"])
            movements.append(Movement(number, "", f"{given} {surname}".strip(), given, surname, _text(row["Riferimento al TUIR"]), _text(row.get("Tipologia di servizio")), row["Valore"], row["Data di assegnazione"]))
    elif provider == "E":
        rows = _xlsx_rows(data, {"Codice Fiscale Dipendente", "Articolo", "Totale", "Data"})
        for number, row in rows:
            movements.append(Movement(number, _text(row["Codice Fiscale Dipendente"]), f"{_text(row['Nome'])} {_text(row['Cognome'])}".strip(), "", "", _text(row["Articolo"]), "", row["Totale"], row["Data"]))
    else:
        raise ConversionError("Provider non riconosciuto.")
    if not movements:
        raise ConversionError("Il file non contiene movimenti da convertire.")
    return movements


def _amount_cents(value: object) -> int:
    if isinstance(value, (int, float, Decimal)):
        raw = str(value)
    else:
        raw = _text(value).replace("€", "").replace("\u00a0", "").replace(" ", "")
        if not raw:
            raise ValueError("importo mancante")
        if "," in raw and "." in raw:
            decimal_separator = "," if raw.rfind(",") > raw.rfind(".") else "."
            thousands_separator = "." if decimal_separator == "," else ","
            raw = raw.replace(thousands_separator, "").replace(decimal_separator, ".")
        elif "," in raw:
            raw = raw.replace(",", ".")
    try:
        amount = Decimal(raw)
    except InvalidOperation as exc:
        raise ValueError("importo non numerico") from exc
    if not amount.is_finite() or amount < 0:
        raise ValueError("importo negativo o non valido")
    cents = amount * 100
    if cents != cents.to_integral_value():
        raise ValueError("importo con più di due decimali")
    return int(cents)


class EmployeeLookup:
    def __init__(self, company: str):
        self.by_fiscal_code: dict[str, set[str]] = defaultdict(set)
        self.by_name: dict[str, set[str]] = defaultdict(set)
        self.fiscal_codes_by_employee: dict[str, set[str]] = defaultdict(set)
        self.companies_by_fiscal_code: dict[str, set[str]] = defaultdict(set)
        with (CONFIG_DIR / "Lista_Dipendenti.csv").open(newline="", encoding="utf-8-sig") as file:
            for row in csv.DictReader(file):
                row_company = row["Codice Ditta"].strip()
                code = row["CodAnagraficoLav"].strip()
                fiscal_code = row["CodiceFiscale"].strip().upper()
                self.companies_by_fiscal_code[fiscal_code].add(row_company)
                if row_company != company:
                    continue
                self.by_fiscal_code[fiscal_code].add(code)
                self.by_name[_name_key(row["Nome Dipendente"])].add(code)
                self.fiscal_codes_by_employee[code].add(fiscal_code)
        if not self.by_fiscal_code:
            raise ConversionError(f"Ditta {company} non presente nell'anagrafica fornita.")

    def find(self, movement: Movement) -> str:
        if movement.fiscal_code:
            codes = self.by_fiscal_code.get(movement.fiscal_code.upper(), set())
        else:
            candidates = {
                _name_key(f"{movement.given_name} {movement.surname}"),
                _name_key(f"{movement.surname} {movement.given_name}"),
            }
            codes = set().union(*(self.by_name.get(candidate, set()) for candidate in candidates))
        if not codes:
            raise ValueError("dipendente non trovato per questa ditta")
        if len(codes) != 1:
            raise ValueError("corrispondenza dipendente ambigua")
        return next(iter(codes))

    def cross_company_error(self, movement: Movement, employee: str) -> str | None:
        fiscal_codes = self.fiscal_codes_by_employee[employee]
        if movement.fiscal_code:
            fiscal_code = movement.fiscal_code.upper()
            if fiscal_code not in fiscal_codes:
                return None
        elif len(fiscal_codes) == 1:
            fiscal_code = next(iter(fiscal_codes))
        else:
            return None
        companies = sorted(self.companies_by_fiscal_code[fiscal_code], key=int)
        if len(companies) < 2:
            return None
        label = "due aziende" if len(companies) == 2 else "più aziende"
        return f"Dipendente duplicato in {label} ({','.join(companies)})"


class WelfareLookup:
    def __init__(self):
        self.mapping: dict[tuple[str, str], str] = {}
        self.normalized: dict[tuple[str, str], set[str]] = defaultdict(set)
        with (CONFIG_DIR / "Codici_Welfare_Voci_Payroll.csv").open(newline="", encoding="utf-8-sig") as file:
            for row in csv.DictReader(file):
                key = (_key(row["Trattamento"]), _key(row["Sottocategoria"]))
                code = row["Codice"].strip()
                old = self.mapping.get(key)
                if old and old != code:
                    raise ConversionError(f"Mapping welfare ambiguo: {row['Trattamento']} / {row['Sottocategoria']}")
                self.mapping[key] = code
                self.normalized[(_tax_key(row["Trattamento"]), key[1])].add(code)

    def _normalized_code(self, treatment: str, subcategory: str) -> str | None:
        codes = self.normalized.get((_tax_key(treatment), subcategory), set())
        if len(codes) > 1:
            raise ValueError(f"mapping welfare ambiguo dopo normalizzazione: {treatment}")
        return next(iter(codes)) if codes else None

    def find(self, movement: Movement) -> str:
        treatment, subcategory = _key(movement.treatment), _key(movement.subcategory)
        if (treatment, subcategory) in self.mapping:
            return self.mapping[(treatment, subcategory)]
        if subcategory:
            code = self._normalized_code(movement.treatment, subcategory)
            if code:
                return code
        if (treatment, "") in self.mapping:
            return self.mapping[(treatment, "")]
        code = self._normalized_code(movement.treatment, "")
        if code:
            return code

        # These provider labels are absent from the supplied mapping. Use the
        # generic treatment only when its meaning is clear from the label.
        plain = re.sub(r"[^a-z0-9]+", "", treatment)
        if "art100" in plain:
            return "370"
        if "art51" in plain and "fbis" in plain:
            return "373"
        if "art51" in plain and "dbis" in plain:
            return "377"
        if "art51" in plain and "c3" in plain:
            return "371"
        if "art51" in plain and "c2" in plain and plain.endswith("f"):
            return "370"
        raise ValueError(f"voce welfare non mappata: {movement.treatment or '(vuota)'}")


def _record(company: str, employee: str, voice: str, cents: int, period: str) -> str:
    if not employee.isascii() or not employee.isdigit() or len(employee) > 6:
        raise ConversionError("Codice dipendente non rappresentabile nel tracciato.")
    if not voice.isascii() or len(voice) > 4:
        raise ConversionError("Codice voce non rappresentabile nel tracciato.")
    if cents > 999_999_999:
        raise ConversionError("Importo aggregato oltre il limite del tracciato.")
    line = f"{int(company):06d}{employee:<6}{voice:<4}{0:07d}{cents:09d}{period}"
    assert len(line) == 38 and line.isascii()
    return line


def convert(provider: str, company: str, period: str, filename: str, data: bytes) -> ConversionResult:
    provider = provider.strip().upper()
    company, period = company.strip(), period.strip()
    if provider not in PROVIDER_EXTENSIONS:
        raise ConversionError("Seleziona un provider valido.")
    if not re.fullmatch(r"[0-9]{1,6}", company):
        raise ConversionError("Il codice ditta deve avere da 1 a 6 cifre.")
    if not re.fullmatch(r"[0-9]{6}", period) or not 1 <= int(period[4:]) <= 12:
        raise ConversionError("Il periodo deve essere nel formato AAAAMM, con mese valido.")
    company = str(int(company))
    if Path(filename).suffix.lower() != PROVIDER_EXTENSIONS[provider]:
        raise ConversionError(f"Per Provider {provider} carica un file {PROVIDER_EXTENSIONS[provider]}.")
    if not data:
        raise ConversionError("Il file caricato è vuoto.")

    employees = EmployeeLookup(company)
    welfare = WelfareLookup()
    movements = _parse_provider(provider, data)
    _validate_period(movements, period)
    totals: dict[tuple[str, str], int] = defaultdict(int)
    issues: list[Issue] = []
    converted_rows = 0
    for movement in movements:
        problems = []
        try:
            employee = employees.find(movement)
        except ValueError as exc:
            problems.append(str(exc))
            employee = None
        if employee is not None:
            duplicate_error = employees.cross_company_error(movement, employee)
            if duplicate_error:
                issues.append(Issue(movement.row, movement.display_name or "—", duplicate_error))
                continue
        try:
            voice = welfare.find(movement)
        except ValueError as exc:
            problems.append(str(exc))
            voice = None
        try:
            cents = _amount_cents(movement.amount)
        except ValueError as exc:
            problems.append(str(exc))
            cents = None
        if problems:
            issues.append(Issue(movement.row, movement.display_name or "—", "; ".join(problems)))
            continue
        totals[(employee, voice)] += cents
        converted_rows += 1

    records = [_record(company, employee, voice, cents, period) for (employee, voice), cents in sorted(totals.items(), key=lambda item: (int(item[0][0]), item[0][1]))]
    content = "\r\n".join(records) + ("\r\n" if records else "")
    return ConversionResult(f"VOCI_{company}_{period}.txt", content, len(movements), converted_rows, len(records), issues)
